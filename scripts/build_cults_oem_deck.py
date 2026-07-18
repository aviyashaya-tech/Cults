from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt


OUTPUT_PATH = Path("/workspace/decks/cults-gummies-oem-hotel-proposal.pptx")
ASSET_DIR = Path("/workspace/assets/cults-brand")
LOGO_LOCKUP = ASSET_DIR / "cults-logo-lockup.png"
HERO_POUCH = ASSET_DIR / "cults-hero-pouch.png"
SECONDARY_POUCH = ASSET_DIR / "cults-secondary-pouch.png"
NEON_BG = ASSET_DIR / "cults-neon-gradient-bg.png"
STAR_ICON = ASSET_DIR / "cults-star-icon.png"
ROOM_LIFESTYLE = ASSET_DIR / "cults-room-lifestyle.png"
STICKER_SHEET = ASSET_DIR / "cults-sticker-sheet.png"
POSTER_MARK = ASSET_DIR / "cults-poster-mark.png"
OPEN_YOUR_EYES_POSTER = ASSET_DIR / "cults-open-your-eyes-poster.png"
SOMETHING_COMING_POSTER = ASSET_DIR / "cults-something-is-coming-poster.png"


BG = RGBColor(8, 8, 12)
CARD = RGBColor(22, 22, 28)
WHITE = RGBColor(245, 245, 245)
MUTED = RGBColor(164, 164, 170)
PINK = RGBColor(240, 45, 140)
LIME = RGBColor(186, 219, 68)
ORANGE = RGBColor(244, 132, 52)
PURPLE = RGBColor(165, 93, 255)


def add_bg(slide, color=BG):
    shape = slide.shapes.add_shape(1, Inches(0), Inches(0), Inches(13.33), Inches(7.5))
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    if NEON_BG.exists():
        slide.shapes.add_picture(str(NEON_BG), Inches(0), Inches(0), Inches(13.33), Inches(7.5))
        overlay = slide.shapes.add_shape(1, Inches(0), Inches(0), Inches(13.33), Inches(7.5))
        overlay.fill.solid()
        overlay.fill.fore_color.rgb = BG
        overlay.fill.transparency = 0.32
        overlay.line.fill.background()


def add_brand_motifs(slide):
    if STAR_ICON.exists():
        slide.shapes.add_picture(str(STAR_ICON), Inches(0.55), Inches(0.8), Inches(0.22), Inches(0.22))
        slide.shapes.add_picture(str(STAR_ICON), Inches(12.55), Inches(0.92), Inches(0.18), Inches(0.18))
        slide.shapes.add_picture(str(STAR_ICON), Inches(12.2), Inches(6.85), Inches(0.18), Inches(0.18))


def add_meta(slide):
    left = slide.shapes.add_textbox(Inches(0.35), Inches(0.15), Inches(3.5), Inches(0.3))
    ltf = left.text_frame
    ltf.clear()
    p = ltf.paragraphs[0]
    p.text = "@drinkcults"
    p.font.name = "Calibri"
    p.font.size = Pt(10)
    p.font.bold = True
    p.font.color.rgb = MUTED

    right = slide.shapes.add_textbox(Inches(9.65), Inches(0.15), Inches(3.3), Inches(0.3))
    rtf = right.text_frame
    rtf.clear()
    p = rtf.paragraphs[0]
    p.text = "Business use  -  Highly confidential"
    p.alignment = PP_ALIGN.RIGHT
    p.font.name = "Calibri"
    p.font.size = Pt(9)
    p.font.color.rgb = MUTED
    add_brand_motifs(slide)


def add_title(slide, title):
    box = slide.shapes.add_textbox(Inches(0.75), Inches(0.72), Inches(12), Inches(0.8))
    tf = box.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    p.text = title
    p.font.name = "Calibri"
    p.font.size = Pt(39)
    p.font.bold = True
    p.font.color.rgb = PINK


def add_body_text(slide, x, y, w, h, lines, size=16, color=WHITE, spacing=10):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    for idx, line in enumerate(lines):
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        p.text = line
        p.level = 0
        p.font.name = "Calibri"
        p.font.size = Pt(size)
        p.font.color.rgb = color
        p.space_after = Pt(spacing)


def add_card(slide, x, y, w, h, title, items, border=PINK):
    card = slide.shapes.add_shape(1, Inches(x), Inches(y), Inches(w), Inches(h))
    card.fill.solid()
    card.fill.fore_color.rgb = CARD
    card.line.color.rgb = border

    t = slide.shapes.add_textbox(Inches(x + 0.25), Inches(y + 0.2), Inches(w - 0.5), Inches(0.5))
    tf = t.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    p.text = title
    p.font.name = "Calibri"
    p.font.size = Pt(20)
    p.font.bold = True
    p.font.color.rgb = border

    add_body_text(slide, x + 0.25, y + 0.8, w - 0.5, h - 1.0, items, size=13, color=MUTED, spacing=6)


def cover(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide)
    add_meta(slide)

    if LOGO_LOCKUP.exists():
        slide.shapes.add_picture(str(LOGO_LOCKUP), Inches(0.72), Inches(1.45), Inches(6.2), Inches(2.0))
    else:
        add_body_text(
            slide,
            0.78,
            1.85,
            8.5,
            1.2,
            ["CULTS"],
            size=64,
            color=PINK,
            spacing=0,
        )
        add_body_text(
            slide,
            0.8,
            3.05,
            6,
            0.6,
            ["ADAPTOGEN GUMMIES"],
            size=24,
            color=LIME,
            spacing=0,
        )

    add_body_text(
        slide,
        0.8,
        3.45,
        8.4,
        0.6,
        ["Recovery Ritual  -  OEM partnership proposal for boutique and luxury hotels"],
        size=14,
        color=MUTED,
        spacing=0,
    )

    if HERO_POUCH.exists():
        slide.shapes.add_picture(str(HERO_POUCH), Inches(8.05), Inches(0.85), Inches(4.85), Inches(6.4))

    if STICKER_SHEET.exists():
        sticker_frame = slide.shapes.add_shape(1, Inches(0.82), Inches(4.96), Inches(2.6), Inches(2.0))
        sticker_frame.fill.solid()
        sticker_frame.fill.fore_color.rgb = CARD
        sticker_frame.line.color.rgb = PINK
        slide.shapes.add_picture(str(STICKER_SHEET), Inches(0.9), Inches(5.04), Inches(2.44), Inches(1.84))

    cta = slide.shapes.add_shape(1, Inches(0.82), Inches(4.4), Inches(4.2), Inches(0.55))
    cta.fill.solid()
    cta.fill.fore_color.rgb = BG
    cta.line.color.rgb = ORANGE
    cta_text = cta.text_frame
    cta_text.clear()
    p = cta_text.paragraphs[0]
    p.text = "RECOVER LIKE YOU MEAN IT"
    p.alignment = PP_ALIGN.CENTER
    p.font.name = "Calibri"
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = ORANGE


def brand_world(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide)
    add_meta(slide)
    add_title(slide, "Brand world")

    add_body_text(
        slide,
        0.9,
        1.55,
        12.0,
        0.65,
        ["Distinctive retro-psychedelic language that turns an amenity into a memory and a shareable hotel moment."],
        size=14,
        color=WHITE,
        spacing=0,
    )

    if OPEN_YOUR_EYES_POSTER.exists():
        poster = slide.shapes.add_shape(1, Inches(0.9), Inches(2.02), Inches(3.92), Inches(4.98))
        poster.fill.solid()
        poster.fill.fore_color.rgb = CARD
        poster.line.color.rgb = ORANGE
        slide.shapes.add_picture(str(OPEN_YOUR_EYES_POSTER), Inches(1.02), Inches(2.14), Inches(3.68), Inches(4.72))

    if SOMETHING_COMING_POSTER.exists():
        poster2 = slide.shapes.add_shape(1, Inches(5.02), Inches(2.02), Inches(3.92), Inches(4.98))
        poster2.fill.solid()
        poster2.fill.fore_color.rgb = CARD
        poster2.line.color.rgb = PINK
        slide.shapes.add_picture(str(SOMETHING_COMING_POSTER), Inches(5.14), Inches(2.14), Inches(3.68), Inches(4.72))

    if STICKER_SHEET.exists():
        board = slide.shapes.add_shape(1, Inches(9.15), Inches(2.02), Inches(3.22), Inches(4.98))
        board.fill.solid()
        board.fill.fore_color.rgb = CARD
        board.line.color.rgb = LIME
        slide.shapes.add_picture(str(STICKER_SHEET), Inches(9.27), Inches(2.17), Inches(2.98), Inches(4.68))

    add_body_text(
        slide,
        0.9,
        7.02,
        12.1,
        0.3,
        ["Campaign cues: OPEN YOUR EYES / SOMETHING IS COMING / BELIEVERS ONLY, translated into premium in-room ritual storytelling."],
        size=11,
        color=MUTED,
        spacing=0,
    )


def intro(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide)
    add_meta(slide)
    add_title(slide, "Introduction")

    if ROOM_LIFESTYLE.exists():
        slide.shapes.add_picture(str(ROOM_LIFESTYLE), Inches(6.72), Inches(1.72), Inches(5.7), Inches(4.95))
        overlay = slide.shapes.add_shape(1, Inches(6.72), Inches(1.72), Inches(5.7), Inches(4.95))
        overlay.fill.solid()
        overlay.fill.fore_color.rgb = BG
        overlay.fill.transparency = 0.42
        overlay.line.fill.background()

    add_body_text(
        slide,
        0.9,
        1.7,
        5.7,
        4.9,
        [
            "Brief business overview",
            "Cults makes functional gummies for modern travelers who want recovery and deeper sleep without the clinical wellness aesthetic.",
            "",
            "Vision statement",
            "To become the in-room ritual guests trust the moment their trip starts to tax their body and sleep cycle.",
            "",
            "Mission statement",
            "Put one honest, effective pouch in every room that needs a fast, feel-good reset.",
        ],
        size=14,
        color=WHITE,
    )

    add_card(
        slide,
        7.0,
        1.8,
        5.35,
        4.8,
        "What's in the pouch",
        [
            "Deep Sleep blend:",
            "reishi + ashwagandha + magnesium + passionflower + l-theanine",
            "",
            "Recovery blend:",
            "cordyceps + tart cherry + turmeric + mangosteen",
            "",
            "No powders. No capsules. No fake pharma voice.",
            "Just one premium pouch guests can actually enjoy using.",
        ],
        border=PINK,
    )


def who_we_are(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide)
    add_meta(slide)
    add_title(slide, "Who we are")

    if SECONDARY_POUCH.exists():
        slide.shapes.add_picture(str(SECONDARY_POUCH), Inches(0.12), Inches(3.25), Inches(2.25), Inches(3.75))

    add_body_text(
        slide,
        0.9,
        1.7,
        6.2,
        4.9,
        [
            "Cults is a Thai functional gummy brand built for people who think wellness can be real, stylish, and simple.",
            "For hotels, that becomes a guest-facing amenity that helps jet-lagged, over-worked, over-social guests sleep and recover faster.",
            "",
            "Best-fit guest moments:",
            "- First night, post-flight",
            "- Post-gym or spa recovery",
            "- Wind-down after late dinner",
            "- Early call-time, next-day reset",
        ],
        size=14,
        color=WHITE,
    )

    add_card(
        slide,
        6.65,
        1.95,
        3.0,
        2.05,
        "Boutique city hotel",
        ["Amenity drawer, turn-down tray"],
        border=PINK,
    )
    add_card(
        slide,
        9.78,
        1.95,
        2.55,
        2.05,
        "Wellness resort",
        ["Spa gift shop, retreat welcome kit"],
        border=ORANGE,
    )
    add_card(
        slide,
        6.65,
        4.2,
        3.0,
        2.05,
        "Business hotel",
        ["Executive floor minibar"],
        border=LIME,
    )
    add_card(
        slide,
        9.78,
        4.2,
        2.55,
        2.05,
        "Private villa",
        ["Welcome hamper, concierge upsell"],
        border=PURPLE,
    )


def market_opportunity(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide)
    add_meta(slide)
    add_title(slide, "Market opportunity")

    if HERO_POUCH.exists():
        slide.shapes.add_picture(str(HERO_POUCH), Inches(10.95), Inches(1.52), Inches(1.95), Inches(3.15))

    add_body_text(
        slide,
        0.9,
        1.7,
        6.2,
        1.6,
        [
            "Sleep tourism is now a leading driver of premium-stay booking behavior.",
            "Hotels are actively searching for low-friction in-room wellness amenities.",
        ],
        size=14,
        color=WHITE,
    )

    add_card(
        slide,
        0.9,
        3.1,
        6.1,
        3.0,
        "Why Cults fits now",
        [
            "The shelf is crowded with generic minibar items.",
            "Cults gives properties a premium, photogenic recovery touchpoint.",
            "Crafted in Thailand with real ingredients and modern design language.",
        ],
        border=ORANGE,
    )

    add_body_text(
        slide,
        7.35,
        1.9,
        5.3,
        4.6,
        [
            "$369B",
            "global luxury hotel market by 2032",
            "",
            "9%",
            "growth in functional wellness beverage/ritual demand",
            "",
            "23%",
            "of guests say personalized stay amenities increase satisfaction",
            "",
            "Cults gives hotels one SKU that is easy to deploy and easy to talk about.",
        ],
        size=22,
        color=LIME,
        spacing=2,
    )


def target_partners(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide)
    add_meta(slide)
    add_title(slide, "Target hotel partners")

    line = slide.shapes.add_shape(1, Inches(0.95), Inches(1.65), Inches(11.9), Inches(0.04))
    line.fill.solid()
    line.fill.fore_color.rgb = LIME
    line.line.fill.background()

    add_body_text(
        slide,
        0.95,
        1.9,
        5.9,
        4.5,
        [
            "Primary",
            "Boutique and lifestyle hotels in Bangkok, Phuket, Chiang Mai, Koh Samui",
            "",
            "Wellness resorts and spa retreats",
            "Properties already investing in sleep and recovery programming",
            "",
            "Independent luxury properties",
            "Fast decision-making, no heavy procurement layers",
        ],
        size=14,
        color=WHITE,
    )

    add_body_text(
        slide,
        6.95,
        1.9,
        5.2,
        4.5,
        [
            "Secondary",
            "International chain wellness floors and executive tiers",
            "",
            "Private villa and estate management",
            "Welcome hampers, concierge bundles, in-villa recovery packs",
            "",
            "Corporate wellness retreats and off-sites",
            "Bulk gifting and premium room-drop moments",
        ],
        size=14,
        color=WHITE,
    )

    add_card(
        slide,
        0.95,
        5.85,
        11.95,
        1.1,
        "The guest they're serving",
        [
            "Long-haul arrivals fighting jet lag, wellness-forward travelers, busy executives, and social travelers who want a premium reset without effort.",
        ],
        border=PINK,
    )


def value_market(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide)
    add_meta(slide)
    add_title(slide, "Our value to the market")

    add_body_text(slide, 0.95, 1.63, 5.0, 0.5, ["The first true in-room recovery amenity in Thailand"], size=16, color=WHITE)

    circles = [
        (1.0, "Low-cost, high-value", "One pouch, priced like a thoughtful add-on, felt like a premium ritual.", PINK),
        (4.45, "Effortless ritual", "No blending. No prep. Tear, eat, and wind down before lights out.", LIME),
        (7.9, "Real ingredients", "Botanicals, mushrooms, tart cherry, magnesium, and tropical fruit.", ORANGE),
    ]

    for x, head, body, color in circles:
        circ = slide.shapes.add_shape(9, Inches(x), Inches(2.15), Inches(3.2), Inches(3.2))
        circ.fill.solid()
        circ.fill.fore_color.rgb = CARD
        circ.line.color.rgb = color
        circ.line.width = Pt(2.2)

        add_body_text(slide, x + 0.3, 3.2, 2.6, 0.7, [head], size=16, color=WHITE, spacing=2)
        add_body_text(slide, x + 0.28, 4.0, 2.65, 1.2, [body], size=11, color=MUTED, spacing=2)


def value_to_you(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide)
    add_meta(slide)
    add_title(slide, "Our value to you")

    cards = [
        (0.95, "1. Guest experience upgrade", "A recovery ritual that photographs beautifully and genuinely helps guests reset."),
        (4.43, "2. Low cost, high perceived value", "Priced like a premium turn-down treat, felt like a top-tier amenity."),
        (7.91, "3. A story your team can tell", "Positions the property in sleep-tourism and modern wellness without capital-heavy buildout."),
    ]
    borders = [ORANGE, LIME, PINK]
    for idx, (x, title, text) in enumerate(cards):
        add_card(slide, x, 1.75, 3.2, 2.3, title, [text], border=borders[idx])

    add_body_text(
        slide,
        0.98,
        4.45,
        12.0,
        2.0,
        [
            "OEM program",
            "Private label: hotel-branded pouch, welcome-note insert, or co-branded Cults x property edition",
            "Low minimums: batch runs through Thai manufacturing partners",
            "Flexible fulfillment: single-property pilot or multi-property rollout",
        ],
        size=14,
        color=WHITE,
    )


def marketing_support(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide)
    add_meta(slide)
    add_title(slide, "Marketing support")

    add_body_text(
        slide,
        0.95,
        1.75,
        5.9,
        4.9,
        [
            "Where to find Cults",
            "Partner-hotel discovery posts across our channels and booking-page referrals.",
            "",
            "In-room QR, post-stay flywheel",
            "QR on pouch links guests to your property story and our direct channel.",
            "",
            "Always-on content",
            "Founder-led social + creator seeding around real hotel stays.",
        ],
        size=14,
        color=WHITE,
    )

    add_card(
        slide,
        6.7,
        1.9,
        5.6,
        4.6,
        "Marketing approach",
        [
            "High-visibility pilot launch at your property events and wellness moments.",
            "",
            "Acquisition:",
            "social + hospitality media + hotel group intros",
            "",
            "Engagement:",
            "in-room QR and guest-generated content moments",
            "",
            "Repurchase:",
            "direct-to-consumer reorder path with hotel referral credit",
        ],
        border=LIME,
    )


def pilot_rollout(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide)
    add_meta(slide)
    add_title(slide, "Pilot rollout and commercial model")

    add_card(
        slide,
        0.95,
        1.8,
        3.85,
        4.5,
        "Phase 1\n30-day pilot",
        [
            "2-3 room categories",
            "100-300 pouches",
            "Guest feedback + reorder signal",
        ],
        border=PINK,
    )
    add_card(
        slide,
        4.95,
        1.8,
        3.85,
        4.5,
        "Phase 2\n90-day scale",
        [
            "Full floor or full property",
            "Co-branded inserts and QR journey",
            "Monthly optimization check-ins",
        ],
        border=LIME,
    )
    add_card(
        slide,
        8.95,
        1.8,
        3.35,
        4.5,
        "Commercial terms",
        [
            "OEM or co-branded options",
            "Low MOQs via local production",
            "Simple reorder cadence",
            "Protected margin architecture",
        ],
        border=ORANGE,
    )

    add_body_text(
        slide,
        0.95,
        6.45,
        12.0,
        0.7,
        ["Target outcome: one Cults pouch in every room that needs better recovery and deeper sleep."],
        size=15,
        color=WHITE,
        spacing=0,
    )


def closing(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide)
    if ROOM_LIFESTYLE.exists():
        slide.shapes.add_picture(str(ROOM_LIFESTYLE), Inches(0), Inches(0), Inches(13.33), Inches(7.5))
        overlay = slide.shapes.add_shape(1, Inches(0), Inches(0), Inches(13.33), Inches(7.5))
        overlay.fill.solid()
        overlay.fill.fore_color.rgb = BG
        overlay.fill.transparency = 0.55
        overlay.line.fill.background()
    add_meta(slide)

    add_body_text(
        slide,
        1.35,
        2.55,
        10.8,
        1.6,
        ["Be the first hotel to put deep sleep in every room."],
        size=52,
        color=PINK,
        spacing=0,
    )
    add_body_text(
        slide,
        1.4,
        4.18,
        10.5,
        0.6,
        ["Join the sleep-tourism wave with a single, honest pouch guests remember."],
        size=16,
        color=MUTED,
        spacing=0,
    )

    cta = slide.shapes.add_shape(1, Inches(4.05), Inches(5.05), Inches(5.2), Inches(0.65))
    cta.fill.solid()
    cta.fill.fore_color.rgb = BG
    cta.line.color.rgb = LIME
    tf = cta.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    p.text = "RECOVER LIKE YOU MEAN IT  -  CULTS OEM HOTEL PROGRAM"
    p.alignment = PP_ALIGN.CENTER
    p.font.name = "Calibri"
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = LIME


def build_deck():
    prs = Presentation()
    prs.slide_width = Inches(13.33)
    prs.slide_height = Inches(7.5)

    cover(prs)
    brand_world(prs)
    intro(prs)
    who_we_are(prs)
    market_opportunity(prs)
    target_partners(prs)
    value_market(prs)
    value_to_you(prs)
    marketing_support(prs)
    pilot_rollout(prs)
    closing(prs)

    OUTPUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    prs.save(OUTPUT_PATH)


if __name__ == "__main__":
    build_deck()

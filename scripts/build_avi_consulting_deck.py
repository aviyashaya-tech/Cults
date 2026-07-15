from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt


OUTPUT_PATH = Path("/workspace/decks/avi-yashaya-consulting-deck.pptx")


NAVY = RGBColor(18, 35, 63)
GOLD = RGBColor(199, 153, 72)
SLATE = RGBColor(54, 69, 89)
LIGHT_BG = RGBColor(247, 249, 252)
WHITE = RGBColor(255, 255, 255)


def add_full_background(slide, color):
    shape = slide.shapes.add_shape(1, Inches(0), Inches(0), Inches(13.33), Inches(7.5))
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()


def add_header_bar(slide, label):
    bar = slide.shapes.add_shape(1, Inches(0), Inches(0), Inches(13.33), Inches(0.55))
    bar.fill.solid()
    bar.fill.fore_color.rgb = NAVY
    bar.line.fill.background()

    tf = bar.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    run = p.add_run()
    run.text = label
    run.font.name = "Calibri"
    run.font.size = Pt(16)
    run.font.bold = True
    run.font.color.rgb = WHITE
    p.alignment = PP_ALIGN.LEFT
    tf.margin_left = Inches(0.25)


def add_title_slide(prs, title, subtitle, contact):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_full_background(slide, NAVY)

    accent = slide.shapes.add_shape(1, Inches(0), Inches(6.8), Inches(13.33), Inches(0.7))
    accent.fill.solid()
    accent.fill.fore_color.rgb = GOLD
    accent.line.fill.background()

    title_box = slide.shapes.add_textbox(Inches(0.9), Inches(1.2), Inches(11.5), Inches(1.5))
    tf = title_box.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    run = p.add_run()
    run.text = title
    run.font.name = "Calibri"
    run.font.size = Pt(44)
    run.font.bold = True
    run.font.color.rgb = WHITE

    sub_box = slide.shapes.add_textbox(Inches(0.9), Inches(2.9), Inches(11.2), Inches(1.7))
    sub_tf = sub_box.text_frame
    sub_tf.word_wrap = True
    sub_tf.clear()
    p = sub_tf.paragraphs[0]
    r = p.add_run()
    r.text = subtitle
    r.font.name = "Calibri"
    r.font.size = Pt(22)
    r.font.color.rgb = RGBColor(223, 230, 239)

    contact_box = slide.shapes.add_textbox(Inches(0.9), Inches(6.9), Inches(11.8), Inches(0.5))
    ctf = contact_box.text_frame
    ctf.clear()
    p = ctf.paragraphs[0]
    rr = p.add_run()
    rr.text = contact
    rr.font.name = "Calibri"
    rr.font.size = Pt(16)
    rr.font.bold = True
    rr.font.color.rgb = NAVY


def add_bullet_slide(prs, header, title, bullets, footer=None):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_full_background(slide, LIGHT_BG)
    add_header_bar(slide, header)

    title_box = slide.shapes.add_textbox(Inches(0.7), Inches(0.9), Inches(12.1), Inches(0.8))
    tf = title_box.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    run = p.add_run()
    run.text = title
    run.font.name = "Calibri"
    run.font.size = Pt(34)
    run.font.bold = True
    run.font.color.rgb = NAVY

    body = slide.shapes.add_textbox(Inches(0.95), Inches(1.8), Inches(11.8), Inches(4.9))
    btf = body.text_frame
    btf.word_wrap = True
    btf.clear()
    for idx, item in enumerate(bullets):
        paragraph = btf.paragraphs[0] if idx == 0 else btf.add_paragraph()
        paragraph.level = 0
        paragraph.text = item
        paragraph.font.name = "Calibri"
        paragraph.font.size = Pt(24 if idx == 0 else 20)
        paragraph.font.color.rgb = SLATE
        paragraph.space_after = Pt(14)

    if footer:
        line = slide.shapes.add_shape(1, Inches(0.95), Inches(6.6), Inches(11.8), Inches(0.03))
        line.fill.solid()
        line.fill.fore_color.rgb = GOLD
        line.line.fill.background()

        footer_box = slide.shapes.add_textbox(Inches(0.95), Inches(6.72), Inches(11.8), Inches(0.6))
        ftf = footer_box.text_frame
        ftf.clear()
        p = ftf.paragraphs[0]
        r = p.add_run()
        r.text = footer
        r.font.name = "Calibri"
        r.font.size = Pt(14)
        r.font.color.rgb = SLATE


def add_two_column_slide(prs, header, title, left_title, left_points, right_title, right_points):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_full_background(slide, LIGHT_BG)
    add_header_bar(slide, header)

    title_box = slide.shapes.add_textbox(Inches(0.7), Inches(0.9), Inches(12.1), Inches(0.8))
    tf = title_box.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    run = p.add_run()
    run.text = title
    run.font.name = "Calibri"
    run.font.size = Pt(30)
    run.font.bold = True
    run.font.color.rgb = NAVY

    left_panel = slide.shapes.add_shape(1, Inches(0.7), Inches(1.9), Inches(5.95), Inches(5.2))
    left_panel.fill.solid()
    left_panel.fill.fore_color.rgb = WHITE
    left_panel.line.color.rgb = RGBColor(216, 223, 232)

    right_panel = slide.shapes.add_shape(1, Inches(6.7), Inches(1.9), Inches(5.95), Inches(5.2))
    right_panel.fill.solid()
    right_panel.fill.fore_color.rgb = WHITE
    right_panel.line.color.rgb = RGBColor(216, 223, 232)

    for panel_x, section_title, points in [
        (0.95, left_title, left_points),
        (6.95, right_title, right_points),
    ]:
        title_shape = slide.shapes.add_textbox(Inches(panel_x), Inches(2.15), Inches(5.45), Inches(0.5))
        stf = title_shape.text_frame
        stf.clear()
        p = stf.paragraphs[0]
        r = p.add_run()
        r.text = section_title
        r.font.name = "Calibri"
        r.font.size = Pt(20)
        r.font.bold = True
        r.font.color.rgb = NAVY

        box = slide.shapes.add_textbox(Inches(panel_x), Inches(2.75), Inches(5.35), Inches(4.1))
        btf = box.text_frame
        btf.clear()
        btf.word_wrap = True
        for idx, point in enumerate(points):
            paragraph = btf.paragraphs[0] if idx == 0 else btf.add_paragraph()
            paragraph.text = point
            paragraph.level = 0
            paragraph.font.name = "Calibri"
            paragraph.font.size = Pt(16)
            paragraph.font.color.rgb = SLATE
            paragraph.space_after = Pt(10)


def build_deck():
    prs = Presentation()
    prs.slide_width = Inches(13.33)
    prs.slide_height = Inches(7.5)

    add_title_slide(
        prs,
        "Avi Yashaya",
        "Partner & CEO, Mahanakhon Beverages Co., Ltd\nNPD Consultant for Beverage Growth Across Southeast Asia\nBuilt an operating platform behind products delivering $2M+ annual sell-through",
        "aviyashaya@gmail.com",
    )

    add_bullet_slide(
        prs,
        "Market Reality",
        "The Double-Edged Dilemma for Beverage Producers in Southeast Asia",
        [
            "SMEs with ambitions of building Southeast Asia's next big beverage brand collide with industrial MOQs and harsh logistical, supply chain, and cash flow crises.",
            "Institutional players struggle to keep up with fast consumer shifts because their data is often past its best-before date by the end of a slow development timeline.",
            "Result: one side cannot scale efficiently, the other cannot adapt fast enough, and both lose speed, margin, and market relevance.",
            "This is the exact operating gap I help clients close.",
        ],
    )

    add_bullet_slide(
        prs,
        "Action Plan",
        "Actionable System I Implement (Built from 10+ Years in Market)",
        [
            "Step 1 - Partner stack: leverage the region's strongest strategic partners for R&D, production, packaging, and route to market.",
            "Step 2 - MOQ-aware build plan: sequence pilot-to-scale production in Thailand and Vietnam across Brewery, Distillery, and RTD to avoid cash burn.",
            "Step 3 - Supply continuity by design: lock critical raw material flow, including hops from USA and New Zealand, with quality/spec alignment and dual-lane planning.",
            "Step 4 - Commercial math: build channel-specific price architecture and margin logic before launch, not after discounting pressure starts.",
            "Step 5 - Execution cadence: run launch governance that shortens development-to-shelf cycles and keeps decisions tied to live market signals.",
        ],
        footer="Let's leverage the region's best strategic partners and build systems that bring Asia's next big beverage to life.",
    )

    add_bullet_slide(
        prs,
        "Global Supply Chain",
        "Hyper-Specific Operating Model: Source Global, Produce Local, Sell Regional",
        [
            "Raw materials: hop procurement from the USA and New Zealand with quality/spec alignment and cost controls.",
            "Production backbone: partner plants in Vietnam and Thailand across Brewery, Distillery, and RTD categories.",
            "Packaging network: coordinated can, bottle, closure, and label suppliers to protect lead time and cash flow.",
            "Regional flow: finished goods exported to convenience store channels in Taiwan and Japan.",
            "Market expansion support: launch blueprints for clients opening new beverage markets in Laos.",
        ],
        footer="Result: diversified supply continuity with practical production optionality across countries and categories.",
    )

    add_two_column_slide(
        prs,
        "Regional Commercialization",
        "Route-to-Market and Pricing Strategy by Market",
        "Thailand + Vietnam",
        [
            "Position portfolio by channel: modern trade, convenience, and on-trade activation.",
            "Set margin architecture for producers, importers/distributors, and retailers.",
            "Build launch calendars tied to local demand windows, events, and consumer behavior.",
            "Run SKU-by-SKU pricing ladders to protect both velocity and gross margin.",
        ],
        "Laos Market Entry",
        [
            "Define import or local production model based on tax and duty realities.",
            "Build first-wave assortment with clear opening price points and premium anchors.",
            "Train distributor teams on product story, outlet pitch, and sell-in mechanics.",
            "Stage rollout by city and outlet class to maximize conversion with controlled risk.",
        ],
    )

    add_two_column_slide(
        prs,
        "Relationships",
        "A Working Network Built Over a Decade",
        "Supply + Production Partners",
        [
            "Hop and ingredient suppliers in the USA and New Zealand.",
            "Manufacturing partners in Thailand and Vietnam covering beer, spirits, and RTD.",
            "Packaging suppliers for cans, bottles, cartons, and labels with scalable capacity.",
            "Quality and operational alignment across partners to speed launch readiness.",
        ],
        "Commercial Partners",
        [
            "Distributors and route-to-market operators across Thailand, Laos, and Vietnam.",
            "Retail and convenience channel relationships in Taiwan and Japan export routes.",
            "On-trade account development through practical pricing and menu placement strategy.",
            "Cross-functional partner management to maintain execution standards and timelines.",
        ],
    )

    add_bullet_slide(
        prs,
        "Category Scope",
        "Cross-Category Execution Capabilities",
        [
            "Brewery: concept-to-commercialization support for craft and premium beer portfolios.",
            "Distillery: liquid positioning, production planning, and go-to-market design for spirits.",
            "RTD: speed-to-market product development and scalable production transition support.",
            "Integrated portfolio strategy to prevent channel conflict and maximize total account value.",
        ],
    )

    add_bullet_slide(
        prs,
        "Commercial Evidence",
        "Selected Outcomes Delivered",
        [
            "Built and operated systems behind Mahanakhon products now exceeding $2M annual retail + on-trade revenue.",
            "Structured practical supply chain pathways linking global ingredient sourcing to regional production execution.",
            "Developed export-ready operating flows for convenience store channels in Taiwan and Japan.",
            "Supported market-opening strategy for Laos with channel-specific launch sequencing and price architecture.",
        ],
        footer="All outcomes based on direct operating involvement in sourcing, production, and commercialization.",
    )

    add_two_column_slide(
        prs,
        "Engagement Model",
        "How I Work With Beverage Companies",
        "Advisory + Operating Sprint",
        [
            "30-90 day focused engagements for market entry, product launch, or supply chain redesign.",
            "Workstreams: sourcing, plant partner selection, packaging readiness, pricing, and channel plan.",
            "Weekly decision cadence with leadership team and execution owners.",
            "Output: clear operating playbook with accountability and rollout checkpoints.",
        ],
        "Fractional NPD Leadership",
        [
            "Embedded support for teams that need senior NPD and commercialization leadership.",
            "Cross-functional bridge between procurement, operations, marketing, and sales.",
            "Focus on profitable scale, not just innovation theater.",
            "Output: repeatable launch process that compounds over multiple product cycles.",
        ],
    )

    add_bullet_slide(
        prs,
        "Ideal Clients",
        "Best Fit Engagements",
        [
            "Global or regional beverage companies entering Southeast Asia.",
            "Established players needing stronger execution in Thailand, Vietnam, or Laos.",
            "Teams expanding across beer, spirits, and RTD categories with shared infrastructure.",
            "Businesses seeking local operator credibility plus strategic clarity at executive level.",
        ],
    )

    add_bullet_slide(
        prs,
        "Contact",
        "Let’s Build the Next Profitable Beverage Platform",
        [
            "Avi Yashaya",
            "Partner & CEO, Mahanakhon Beverages Co., Ltd",
            "NPD Consultant",
            "Email: aviyashaya@gmail.com",
        ],
        footer="Let's leverage the region's best strategic partners for R&D, production, and route to market to bring Asia's next big beverage to life.",
    )

    OUTPUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    prs.save(OUTPUT_PATH)


if __name__ == "__main__":
    build_deck()

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt


OUTPUT_PATH = Path("/workspace/decks/avi-yashaya-carabao-strategy.pptx")
ASSET_DIR = Path("/workspace/assets/carabao-project/normalized")

CARABAO_LOGO = ASSET_DIR / "carabao_logo.png"
SINGHA_LOGO = ASSET_DIR / "singha_logo.png"
TAWANDANG_LOGO = ASSET_DIR / "tawandang_logo.png"

CARABAO_LAGER = ASSET_DIR / "carabao_lager.png"
CHANG_CLASSIC = ASSET_DIR / "chang_classic.png"
LEO_BEER = ASSET_DIR / "leo_beer.png"
CHAO_SUNGTHONG = ASSET_DIR / "chao_sungthong.png"
MAHANAKHON_WHITE_ALE = ASSET_DIR / "mahanakhon_white_ale.png"
LEO_SUPREME = ASSET_DIR / "leo_supreme.png"
SINGHA_RESERVE = ASSET_DIR / "singha_reserve.png"


BG = RGBColor(12, 18, 32)
CARD = RGBColor(23, 33, 53)
CARD_ALT = RGBColor(34, 44, 66)
WHITE = RGBColor(245, 248, 255)
MUTED = RGBColor(171, 185, 213)
ACCENT_GOLD = RGBColor(245, 176, 65)
ACCENT_GREEN = RGBColor(94, 201, 136)
ACCENT_RED = RGBColor(235, 87, 87)


def add_bg(slide, color=BG):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.33), Inches(7.5))
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()


def add_meta(slide, label):
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.33), Inches(0.44))
    bar.fill.solid()
    bar.fill.fore_color.rgb = RGBColor(18, 27, 46)
    bar.line.fill.background()

    tf = bar.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    p.text = label
    p.alignment = PP_ALIGN.LEFT
    p.font.name = "Calibri"
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = MUTED
    tf.margin_left = Inches(0.3)


def add_title(slide, title, subtitle=None):
    box = slide.shapes.add_textbox(Inches(0.72), Inches(0.62), Inches(8.9), Inches(1.3))
    tf = box.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    p.text = title
    p.font.name = "Calibri"
    p.font.size = Pt(36)
    p.font.bold = True
    p.font.color.rgb = WHITE

    if subtitle:
        sub = slide.shapes.add_textbox(Inches(0.72), Inches(1.52), Inches(11.6), Inches(0.72))
        stf = sub.text_frame
        stf.clear()
        sp = stf.paragraphs[0]
        sp.text = subtitle
        sp.font.name = "Calibri"
        sp.font.size = Pt(14)
        sp.font.color.rgb = MUTED


def add_bullets(slide, x, y, w, h, bullets, size=17, color=WHITE, space_after=8):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.clear()
    for idx, bullet in enumerate(bullets):
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        p.text = bullet
        p.level = 0
        p.font.name = "Calibri"
        p.font.size = Pt(size)
        p.font.color.rgb = color
        p.space_after = Pt(space_after)


def add_card(slide, x, y, w, h, title, subtitle=None, fill=CARD):
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    card.fill.solid()
    card.fill.fore_color.rgb = fill
    card.line.color.rgb = RGBColor(48, 61, 90)

    tf = card.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    p.text = title
    p.font.name = "Calibri"
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = WHITE
    p.space_after = Pt(3)

    if subtitle:
        p2 = tf.add_paragraph()
        p2.text = subtitle
        p2.font.name = "Calibri"
        p2.font.size = Pt(11)
        p2.font.color.rgb = MUTED
    return card


def add_image_or_placeholder(slide, path, x, y, w, h, label):
    if path.exists():
        slide.shapes.add_picture(str(path), Inches(x), Inches(y), Inches(w), Inches(h))
        return

    ph = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    ph.fill.solid()
    ph.fill.fore_color.rgb = CARD_ALT
    ph.line.color.rgb = RGBColor(64, 80, 113)
    tf = ph.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    p.text = label
    p.font.name = "Calibri"
    p.font.size = Pt(10)
    p.font.color.rgb = MUTED
    p.alignment = PP_ALIGN.CENTER


def cover(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide)
    add_meta(slide, "CARABAO BREWERY | LONG-TERM STRATEGY | AVI YASHAYA")

    glow = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(6.7), Inches(13.33), Inches(0.8))
    glow.fill.solid()
    glow.fill.fore_color.rgb = ACCENT_GOLD
    glow.line.fill.background()

    add_title(
        slide,
        "Get Brand Identity On Par With Operational Modernization",
        "A focused growth blueprint for Carabao Brewery built by Avi Yashaya, Partner & CEO of Mahanakhon Beverages.",
    )

    add_bullets(
        slide,
        0.76,
        2.35,
        7.1,
        2.8,
        [
            "10+ years building beverage systems across Thailand, Vietnam, and Laos.",
            "Global sourcing + regional production model proven in beer, spirits, and RTD.",
            "Direct track record: products generating $2M+ annual retail and on-trade sell-through.",
            "Goal: reposition Carabao beer from volume play to modern relevance and profitable share gain.",
        ],
        size=17,
    )

    if CARABAO_LOGO.exists():
        logo_card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(8.2), Inches(1.0), Inches(4.4), Inches(2.2))
        logo_card.fill.solid()
        logo_card.fill.fore_color.rgb = CARD
        logo_card.line.color.rgb = RGBColor(67, 80, 111)
        slide.shapes.add_picture(str(CARABAO_LOGO), Inches(8.55), Inches(1.2), Inches(3.7), Inches(1.8))

    add_image_or_placeholder(slide, CARABAO_LAGER, 8.35, 3.58, 1.25, 1.8, "Carabao Lager")
    add_image_or_placeholder(slide, LEO_SUPREME, 9.95, 3.58, 1.25, 1.8, "Leo Supreme")
    add_image_or_placeholder(slide, SINGHA_RESERVE, 11.55, 3.58, 1.25, 1.8, "Singha Reserve")

    contact = slide.shapes.add_textbox(Inches(0.78), Inches(6.86), Inches(10.0), Inches(0.5))
    ctf = contact.text_frame
    ctf.clear()
    cp = ctf.paragraphs[0]
    cp.text = "Avi Yashaya | aviyashaya@gmail.com"
    cp.font.name = "Calibri"
    cp.font.size = Pt(14)
    cp.font.bold = True
    cp.font.color.rgb = BG


def dilemma(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide)
    add_meta(slide, "PROBLEM HOOK")
    add_title(slide, "The Double-Edged Dilemma in Thai Beer")

    add_card(
        slide,
        0.75,
        1.8,
        5.95,
        4.95,
        "SME ambition vs. industrial reality",
        "High MOQs, complex logistics, and tight cash cycles punish innovation speed.",
        fill=CARD,
    )
    add_card(
        slide,
        6.93,
        1.8,
        5.65,
        4.95,
        "Incumbent scale vs. cultural velocity",
        "Large players optimize plants but often launch too slowly for new urban demand patterns.",
        fill=CARD_ALT,
    )

    add_bullets(
        slide,
        1.02,
        2.7,
        5.35,
        3.6,
        [
            "Growth is constrained by production economics before consumer demand is validated.",
            "Commercial teams discount to move volume instead of improving brand value.",
            "Brand architecture remains traditional while young urban drinkers move to lifestyle-driven choices.",
        ],
        size=14,
        color=MUTED,
        space_after=9,
    )
    add_bullets(
        slide,
        7.2,
        2.7,
        5.0,
        3.6,
        [
            "Data arrives late by the time products reach market.",
            "Core identity can become functional and price-led rather than aspirational.",
            "Result: SAM is under-penetrated and SOM growth stalls.",
        ],
        size=14,
        color=MUTED,
        space_after=9,
    )


def brand_frames(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide)
    add_meta(slide, "FIVE BRANDING FRAMES")
    add_title(slide, "Thai Beer Positioning Is Not One Market, It Is Five Frames")

    add_card(slide, 0.75, 1.95, 2.35, 2.4, "1) Animal / Power", "Chang, Leo, Carabao")
    add_card(slide, 3.23, 1.95, 2.35, 2.4, "2) Myth / Heritage", "Singha, Chao Sungthong")
    add_card(slide, 5.71, 1.95, 2.35, 2.4, "3) Geographic", "Mahanakhon, Chiang Mai")
    add_card(slide, 8.19, 1.95, 2.35, 2.4, "4) Subculture", "Outlaw craft, gatekept niches")
    add_card(slide, 10.67, 1.95, 2.0, 2.4, "5) Modern Aspirational", "The open growth gap", fill=RGBColor(33, 58, 46))

    add_image_or_placeholder(slide, CHANG_CLASSIC, 1.02, 2.85, 0.58, 0.78, "Chang")
    add_image_or_placeholder(slide, LEO_BEER, 1.67, 2.85, 0.58, 0.78, "Leo")
    add_image_or_placeholder(slide, CARABAO_LAGER, 2.33, 2.85, 0.58, 0.78, "Carabao")

    add_image_or_placeholder(slide, SINGHA_LOGO, 3.5, 2.72, 0.8, 0.95, "Singha")
    add_image_or_placeholder(slide, CHAO_SUNGTHONG, 4.45, 2.82, 0.72, 0.92, "Chao")

    add_image_or_placeholder(slide, MAHANAKHON_WHITE_ALE, 6.45, 2.82, 0.85, 0.86, "Mahanakhon")

    add_bullets(
        slide,
        0.82,
        4.65,
        12.2,
        2.2,
        [
            "Masscult/traditional frames still dominate spend, but growth is shifting toward culturally fluent lifestyle positioning.",
            "Frame 5 is where younger urban consumers reward style, social signaling, and daily-luxury accessibility.",
        ],
        size=14,
        color=MUTED,
    )


def limitations(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide)
    add_meta(slide, "CURRENT LIMITATIONS")
    add_title(slide, "Where Current Framing Limits Growth Velocity")

    add_card(
        slide,
        0.78,
        1.9,
        4.05,
        4.9,
        "Carabao",
        "Strong awareness, but still coded by many younger urbanites as 'heavy worker' or 'cheap night out'.",
    )
    add_card(
        slide,
        4.98,
        1.9,
        4.05,
        4.9,
        "Tawandang",
        "High familiarity, but can feel rigid or 'what my dad drinks' in modern casual lifestyle spaces.",
    )
    add_card(
        slide,
        9.18,
        1.9,
        3.4,
        4.9,
        "Pattaya / Phuket / CNX labels",
        "Local relevance can trap brands in a tourist-frame and reduce national scalability.",
    )

    add_image_or_placeholder(slide, CARABAO_LOGO, 1.35, 3.5, 2.8, 1.15, "Carabao logo")
    add_image_or_placeholder(slide, TAWANDANG_LOGO, 5.45, 3.5, 3.1, 1.15, "Tawandang logo")
    add_image_or_placeholder(slide, LEO_BEER, 9.7, 3.3, 2.4, 1.55, "Category reference")


def sam_som(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide)
    add_meta(slide, "SAM TO SOM")
    add_title(slide, "Carabao Beer Presentation Is Not Optimized for SAM")

    add_card(
        slide,
        0.8,
        2.0,
        4.1,
        4.6,
        "Current state",
        "Strong distribution logic, weaker lifestyle translation in premium-adjacent urban occasions.",
    )
    add_card(
        slide,
        4.98,
        2.0,
        4.1,
        4.6,
        "Unlock SAM",
        "Reframe identity into modern-retro aspirational cues without losing Thai mass resonance.",
        fill=RGBColor(29, 55, 45),
    )
    add_card(
        slide,
        9.16,
        2.0,
        3.45,
        4.6,
        "Increase SOM",
        "Product and pack upgrades convert broader addressable audience into repeat buyers.",
        fill=RGBColor(51, 36, 33),
    )

    add_bullets(
        slide,
        1.08,
        3.0,
        3.6,
        3.2,
        ["Price-led", "Segment-coded", "Limited aspiration cues"],
        size=14,
        color=MUTED,
    )
    add_bullets(
        slide,
        5.25,
        3.0,
        3.6,
        3.2,
        ["Modernized visual language", "Occasion-specific proposition", "Urban social proof systems"],
        size=14,
        color=MUTED,
    )
    add_bullets(
        slide,
        9.45,
        3.0,
        2.9,
        3.2,
        ["Higher trade-up rate", "Better margin mix", "Higher share in high-value outlets"],
        size=14,
        color=MUTED,
    )


def boonrawd_case(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide)
    add_meta(slide, "PROOF OF PRINCIPLE")
    add_title(slide, "Boon Rawd Shows the Playbook Works")

    add_card(
        slide,
        0.8,
        1.95,
        6.05,
        4.95,
        "Example: Leo Supreme + Singha Reserve",
        "Classic brand equity extended into premiumized, modernized sub-lines.",
        fill=CARD,
    )
    add_card(
        slide,
        7.0,
        1.95,
        5.55,
        4.95,
        "Strategic implication for Carabao",
        "Retain mass-scale strengths while building a modern aspirational tier to expand profitable share.",
        fill=CARD_ALT,
    )

    add_image_or_placeholder(slide, LEO_SUPREME, 1.25, 3.0, 2.35, 2.35, "Leo Supreme")
    add_image_or_placeholder(slide, SINGHA_RESERVE, 3.9, 3.02, 2.52, 2.3, "Singha Reserve")

    add_bullets(
        slide,
        7.25,
        3.0,
        5.05,
        3.4,
        [
            "Use premium line extensions to reset brand perception without abandoning core volume.",
            "Anchor identity in clearer lifestyle codes (not only traditional category cues).",
            "Build faster innovation loops tied to channel-specific consumer behavior.",
        ],
        size=14,
        color=MUTED,
    )


def mahanakhon_case(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide)
    add_meta(slide, "RELEVANT CREDIBILITY")
    add_title(slide, "Mahanakhon: From Geographic Niche to Aspirational Relevance")

    add_image_or_placeholder(slide, MAHANAKHON_WHITE_ALE, 0.92, 2.2, 2.6, 2.9, "Mahanakhon")

    add_bullets(
        slide,
        3.85,
        2.0,
        8.2,
        4.95,
        [
            "We had to solve the same Frame-3 trap: local pride alone was not enough to scale.",
            "Shifted through collaboration, social awareness, and aspirational storytelling.",
            "Linked brand language with commercial systems: sourcing, production, and route-to-market cadence.",
            "Result: sustained product traction and $2M+ annual sell-through across retail and on-trade.",
            "This operating path is directly portable to Carabao's beer reset.",
        ],
        size=16,
        color=WHITE,
        space_after=11,
    )


def operating_plan(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide)
    add_meta(slide, "ACTION FRAMEWORK")
    add_title(slide, "Action Plan: Identity + Product + Commercial System")

    phases = [
        (
            0.8,
            "Phase 1 | Reframe",
            "Rebuild visual and verbal identity around modern Thai confidence, not low-price cues.",
            ACCENT_GOLD,
        ),
        (
            3.95,
            "Phase 2 | Product",
            "Launch targeted line upgrades (liquid profile, pack architecture, and occasion fit).",
            ACCENT_GREEN,
        ),
        (
            7.1,
            "Phase 3 | Route-to-Market",
            "Prioritize outlets where perception and margin uplift can happen fastest.",
            RGBColor(99, 149, 255),
        ),
        (
            10.25,
            "Phase 4 | Scale",
            "Codify playbook and expand once premium-to-core halo effect is measurable.",
            ACCENT_RED,
        ),
    ]

    for x, title, detail, color in phases:
        card = add_card(slide, x, 2.2, 2.85, 4.6, title, detail, fill=CARD)
        card.line.color.rgb = color


def product_dev(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide)
    add_meta(slide, "PRODUCT DEVELOPMENT TO LIFT SOM")
    add_title(slide, "Priority Product Moves to Convert SAM Into SOM")

    add_card(slide, 0.78, 1.95, 4.1, 4.95, "Platform 1: Urban Core", "Cleaner lager profile, upgraded pack design, stronger social occasion fit.")
    add_card(slide, 4.96, 1.95, 4.1, 4.95, "Platform 2: Premium Bridge", "Limited/seasonal variants that cue craft-quality while staying accessible.")
    add_card(slide, 9.14, 1.95, 3.44, 4.95, "Platform 3: Channel Exclusives", "Account-specific SKUs for modern trade, convenience, and on-trade.")

    add_image_or_placeholder(slide, CARABAO_LAGER, 1.55, 3.35, 1.05, 1.42, "Carabao")
    add_image_or_placeholder(slide, LEO_SUPREME, 5.7, 3.35, 1.05, 1.42, "Reference")
    add_image_or_placeholder(slide, SINGHA_RESERVE, 10.0, 3.35, 1.6, 1.38, "Reference")

    add_bullets(
        slide,
        1.0,
        5.05,
        11.9,
        1.7,
        [
            "Design rule: every SKU must carry both volume logic and identity logic.",
            "Commercial rule: sequence innovation by margin upside and outlet influence, not by internal preference.",
        ],
        size=13,
        color=MUTED,
        space_after=7,
    )


def engagement(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide, color=RGBColor(14, 22, 38))
    add_meta(slide, "ENGAGEMENT MODEL")
    add_title(slide, "Execution With Avi Yashaya")

    add_card(
        slide,
        0.9,
        2.0,
        6.15,
        4.6,
        "Scope",
        "Brand architecture, product strategy, partner orchestration, and route-to-market execution cadence.",
        fill=CARD,
    )
    add_card(
        slide,
        7.2,
        2.0,
        5.2,
        4.6,
        "KPIs",
        "Perception lift, premium mix growth, repeat purchase, and share gains in priority urban channels.",
        fill=CARD_ALT,
    )

    add_bullets(
        slide,
        1.25,
        3.05,
        5.7,
        3.0,
        [
            "1) Diagnostic sprint",
            "2) 90-day transformation blueprint",
            "3) Weekly implementation governance",
            "4) Quantified SAM-to-SOM scorecard",
        ],
        size=15,
        color=MUTED,
    )
    add_bullets(
        slide,
        7.55,
        3.05,
        4.8,
        2.7,
        [
            "Outcome target: move from price-coded beer to modern Thai choice architecture.",
            "Contact: aviyashaya@gmail.com",
        ],
        size=15,
        color=WHITE,
    )


def source_notes(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide, color=RGBColor(13, 20, 35))
    add_meta(slide, "VISUAL REFERENCES")
    add_title(slide, "Web-Sourced Logos and Product References Included")

    add_bullets(
        slide,
        0.82,
        2.0,
        12.0,
        4.8,
        [
            "Carabao logo: images.seeklogo.com",
            "Singha logo: Wikimedia Commons (File: Singha_Beer_Logo.png)",
            "Tawandang brand visual: tawandang.com",
            "Product references: Untappd pages for Carabao Lager, Chang Classic, Leo, Singha, Chao Sungthong, Mahanakhon White Ale, Leo Supreme, and Singha Reserve.",
            "Supplemental bottle visual: Wikimedia Commons (File: Thailand-Leo-Beer.jpg).",
        ],
        size=14,
        color=MUTED,
        space_after=10,
    )


def build_deck():
    prs = Presentation()
    prs.slide_width = Inches(13.33)
    prs.slide_height = Inches(7.5)

    cover(prs)
    dilemma(prs)
    brand_frames(prs)
    limitations(prs)
    sam_som(prs)
    boonrawd_case(prs)
    mahanakhon_case(prs)
    operating_plan(prs)
    product_dev(prs)
    engagement(prs)
    source_notes(prs)

    OUTPUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    prs.save(OUTPUT_PATH)


if __name__ == "__main__":
    build_deck()
